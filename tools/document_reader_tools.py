from pathlib import Path
from tools.command_guard import get_workspace, is_inside_workspace

MAX_CHARS = 12000
MAX_SEARCH_RESULTS = 50
MAX_PDF_PAGES = 20
MAX_SHEETS = 3
MAX_SPREADSHEET_ROWS = 30

SEARCHABLE_EXTENSIONS = [
    ".txt", ".md", ".py", ".php", ".js", ".jsx",
    ".css", ".html", ".json", ".csv"
]


# ============================================================
# Shared Helpers
# ============================================================

def _limit_output(text: str, max_chars: int = MAX_CHARS) -> str:
    text = text.strip()
    return text[:max_chars] if text else ""


def _resolve_project_file(file_path: str):
    workspace, error = get_workspace()
    if error:
        return None, error

    target = (workspace / file_path).resolve()

    if not is_inside_workspace(target, workspace):
        return None, "Blocked unsafe path access."

    if not target.exists() or not target.is_file():
        return None, "File not found in current project."

    return target, None


def _safe_read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def _missing_dependency(package_name: str, install_command: str) -> str:
    return (
        f"{package_name} is required for this feature.\n\n"
        f"Install with:\n{install_command}"
    )


# ============================================================
# Phase 186 — Local Document Search
# ============================================================

def local_document_search(query: str) -> str:
    workspace, error = get_workspace()
    if error:
        return error

    query = query.lower().strip()
    if not query:
        return "Search query is required."

    matches = []

    for path in workspace.rglob("*"):
        if not path.is_file():
            continue

        if path.suffix.lower() not in SEARCHABLE_EXTENSIONS:
            continue

        content = _safe_read_text(path)
        relative_path = str(path.relative_to(workspace))

        if query in content.lower() or query in path.name.lower():
            matches.append(relative_path)

        if len(matches) >= MAX_SEARCH_RESULTS:
            break

    if not matches:
        return "No matching documents found."

    return "LOCAL DOCUMENT SEARCH RESULTS\n\n" + "\n".join(
        f"- {match}" for match in matches
    )


# ============================================================
# Phase 187 — PDF Reader
# ============================================================

def _extract_pdf_text(target: Path) -> str:
    try:
        import PyPDF2
    except Exception:
        raise ImportError("PyPDF2")

    text_blocks = []

    with target.open("rb") as file:
        reader = PyPDF2.PdfReader(file)

        for index, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            page_text = page.extract_text() or ""
            text_blocks.append(f"\n--- PAGE {index} ---\n{page_text}")

    return "\n".join(text_blocks)


def read_pdf(file_path: str) -> str:
    target, error = _resolve_project_file(file_path)
    if error:
        return error

    try:
        output = _extract_pdf_text(target)
        output = _limit_output(output)

        return output if output else "No readable PDF text found."

    except ImportError:
        return _missing_dependency(
            "PDF reader",
            "pip install PyPDF2"
        )
    except Exception as e:
        return f"PDF read failed: {e}"


# ============================================================
# Phase 188 — DOCX Reader
# ============================================================

def _extract_docx_text(target: Path) -> str:
    try:
        from docx import Document
    except Exception:
        raise ImportError("python-docx")

    document = Document(str(target))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(paragraphs)


def read_docx(file_path: str) -> str:
    target, error = _resolve_project_file(file_path)
    if error:
        return error

    try:
        output = _extract_docx_text(target)
        output = _limit_output(output)

        return output if output else "No readable DOCX text found."

    except ImportError:
        return _missing_dependency(
            "DOCX reader",
            "pip install python-docx"
        )
    except Exception as e:
        return f"DOCX read failed: {e}"


# ============================================================
# Phase 189 — Spreadsheet Reader
# ============================================================

def _format_spreadsheet_value(value) -> str:
    return str(value) if value is not None else ""


def _extract_spreadsheet_text(target: Path) -> str:
    try:
        import openpyxl
    except Exception:
        raise ImportError("openpyxl")

    workbook = openpyxl.load_workbook(
        str(target),
        read_only=True,
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets[:MAX_SHEETS]:
        lines.append(f"\nSHEET: {sheet.title}")

        for row in sheet.iter_rows(
            max_row=MAX_SPREADSHEET_ROWS,
            values_only=True
        ):
            values = [_format_spreadsheet_value(value) for value in row]
            lines.append(" | ".join(values))

    return "\n".join(lines)


def read_spreadsheet(file_path: str) -> str:
    target, error = _resolve_project_file(file_path)
    if error:
        return error

    try:
        output = _extract_spreadsheet_text(target)
        output = _limit_output(output)

        return output if output else "No readable spreadsheet data found."

    except ImportError:
        return _missing_dependency(
            "Spreadsheet reader",
            "pip install openpyxl"
        )
    except Exception as e:
        return f"Spreadsheet read failed: {e}"


# ============================================================
# Phase 190 — Image OCR
# ============================================================

def _extract_image_text(target: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except Exception:
        raise ImportError("pillow-pytesseract")

    image = Image.open(target)
    return pytesseract.image_to_string(image)


def image_ocr_option(file_path: str) -> str:
    target, error = _resolve_project_file(file_path)
    if error:
        return error

    try:
        output = _extract_image_text(target)
        output = _limit_output(output)

        if not output:
            return "OCR completed, but no readable text was detected."

        return (
            "IMAGE OCR RESULT — PHASE 190\n\n"
            f"File: {target.name}\n\n"
            f"{output}"
        )

    except ImportError:
        return _missing_dependency(
            "Image OCR",
            "pip install pillow pytesseract\nsudo apt install tesseract-ocr -y"
        )
    except Exception as e:
        return f"Image OCR failed: {e}"


# ============================================================
# Phase 191 — Screenshot Understanding Status
# ============================================================

def screenshot_understanding_status() -> str:
    return """SCREENSHOT UNDERSTANDING — PHASE 191

Current Mode:
Planning stub only.

Recommended Safe Build:
1. Capture screenshot only after command confirmation.
2. Store screenshot locally.
3. Run OCR or local vision model.
4. Summarize visible UI.
5. Never click or type automatically without confirmation.
"""


# ============================================================
# Help
# ============================================================

def document_reader_help() -> str:
    return """DOCUMENT READER COMMANDS — PHASES 186–191

186. search documents for <query>
187. read pdf <path>
188. read docx <path>
189. read spreadsheet <path>
190. image ocr option <path>
     read image text <path>
191. screenshot understanding status
"""